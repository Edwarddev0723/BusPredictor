import argparse
import os
import pandas as pd
from GetInfo import BusSechudleInfo, RouteRealTime
from Tool import FindClosestTime, CalDiffTime
from openpyxl import Workbook, load_workbook
from Tool import InitHistoryDriveTimeTable, GetCarID, isValid, GetStasticDriveTime, GetGroundTruth, GetHistoryStayTime, GetIQR, GetRatio, InitKmeansCenter, GetKmeans, PredictedStatus
from PredictedStatus import GetModel

if __name__ == "__main__":
    """
    針對路線、方向、星期幾去整理對應的訓練集
    """
    parser = argparse.ArgumentParser()
    parser.add_argument("--routeid", type=int)
    parser.add_argument("--direction", type=int)
    parser.add_argument("--day", type=str)
    parser.add_argument("--start_date", type=str, default="2025-08-01")
    parser.add_argument("--end_date", type=str, default="2025-10-31")
    opt = parser.parse_args()
    print(opt)

    ROUTEID, DIRECTION, DAY, START_DATE, END_DATE = opt.routeid, opt.direction, opt.day, opt.start_date, opt.end_date
    K = 10 if ROUTEID != 609 else 30

    """
    這部分程式碼跟StatisticDataPrepare.py 檔案是一樣的，
    1. 先抓取起程站資訊
    2. 根據起程站點將時刻表以及路線資訊取出。
    3. 針對不同的星期，取出整理好的dict
    格式:
    {20251001: df1, 20251008: df2, ...}
    """
    df = pd.read_csv("./Info/v_stg_ibus_dailytimetable.csv")
    query = (df["routeid"] == ROUTEID) & (df["direction"] == DIRECTION)
    DEPARTSTOP = df[query]["departurestopname_zh_tw"].unique()[0]

    if int(DAY) <= 5: busSechudleInfo = BusSechudleInfo(ROUTEID, DIRECTION, "2025-07-25 00:00:00.000", DEPARTSTOP)
    else: busSechudleInfo = BusSechudleInfo(ROUTEID, DIRECTION, "2025-07-26 00:00:00.000", DEPARTSTOP)
    TimeTable = busSechudleInfo.GetTimeTable()
    RouteTable = busSechudleInfo.GetRouteInfo()

    RealTimeDFDict = RouteRealTime(ROUTEID, DIRECTION, DAY, START_DATE, END_DATE).GetRealTimeDF()

    """
    建立對應資料夾，用於後續的資料新增
    """
    if not os.path.exists(f"training_dataset/{ROUTEID}/{DIRECTION}"): 
        os.makedirs(f"training_dataset/{ROUTEID}/{DIRECTION}")

    """
    由於先前整理好的歷史行駛時間統計表，還是會有缺值(資料真的有夠破)出現，所以先前方法是應用google api 取出站到站的資訊
    在程式一開始，需要先將該表格初始化好，以便後續的資料使用
    """
    HistoryDriveTimeTable = InitHistoryDriveTimeTable(ROUTEID, DAY, DIRECTION, DEPARTSTOP)
    KmeansGroup1, KmeansGroup2 = InitKmeansCenter(ROUTEID, DAY, DIRECTION)

    """
    初始化一個Excel檔案，用來存放結果
    """
    ResultWB = Workbook()

    for Stop in range(1, len(RouteTable)+1):
        ResultWS = ResultWB.create_sheet(title=f"第{Stop}站")
        ResultWS.append(["Date", "SechudeleIndex", "h_driveTime", "Ratio", "h_stayTime", "c_stayTime", "kmeans1", "kmeans2", "PredictedDriveTime", "GroundTruth"])
        ResultWS.column_dimensions["A"].width = 20
        ResultWS.column_dimensions["B"].width = 20
        ResultWS.column_dimensions["C"].width = 20
        ResultWS.column_dimensions["D"].width = 20
        ResultWS.column_dimensions["E"].width = 20
        ResultWS.column_dimensions["F"].width = 20
        ResultWS.column_dimensions["G"].width = 20
        ResultWS.column_dimensions["H"].width = 20
        ResultWS.column_dimensions["I"].width = 20
        ResultWS.column_dimensions["J"].width = 20

    ResultWB.save(f"training_dataset/{ROUTEID}/{DIRECTION}/dataset_{DAY}.xlsx")
    
    """
    開始進行資料整理, 迴圈順序為日期 ===> 時刻表 ===> 每一個站點
    """
    for DATE in RealTimeDFDict.keys():
        RealTimeDF = RealTimeDFDict[DATE]
        RealTimeDF["gpstime"] = pd.to_datetime(RealTimeDF["gpstime"]).dt.tz_localize(None)
        RealTimeDF["gpstime"] = RealTimeDF["gpstime"].dt.time

        for SechudleIndex, SechudleTime in enumerate(TimeTable):
            carID = GetCarID(RealTimeDF, SechudleTime, ROUTEID, DIRECTION)
            if carID == -1: continue
            PreviousStopDepartTime = f"{SechudleTime}:00"
            """
            將剛剛初始化的exel表格載入
            """
            wb = load_workbook(f"training_dataset/{ROUTEID}/{DIRECTION}/dataset_{DAY}.xlsx")

            """
            維護當前時段，過往所有站點的狀態串列，例如: [1, 2, 2, 3]
            """
            BusStatus = []

            """
            載入該班次的模型，以及預設
            """
            model, device = GetModel(ROUTEID, DIRECTION, DAY, SechudleIndex+1)
            DefaultDF = pd.read_excel("predictions.xlsx")
            Query = DefaultDF["班次"].str.contains(f"{ROUTEID}_{DIRECTION}_{DAY}_第{SechudleIndex+1}班")
            DefaultValue = DefaultDF[Query]["Cluster"].values[0]

            for CurrentStop in range(1, len(RouteTable)):
                """
                取出當前站的等待時間
                1. 如果是第一站 ===> 等待時間為0
                2. 如果當前站沒有進站或是離站資料 ===> 等待時間為0
                """
                if CurrentStop == 1:
                    CurrentStopStayTime = 0
                else:
                    ArrivalQuery = (RealTimeDF["routeid"] == ROUTEID) & (RealTimeDF["direction"] == DIRECTION) & (RealTimeDF["stopsequence"] == CurrentStop) & (RealTimeDF["a2eventtype"] == 1) & (RealTimeDF["platenumb"] == carID)
                    ArrivalResult = RealTimeDF[ArrivalQuery]["gpstime"].tolist()
                    ArrivalResult = [str(time) for time in ArrivalResult]

                    DepartQuery = (RealTimeDF["routeid"] == ROUTEID) & (RealTimeDF["direction"] == DIRECTION) & (RealTimeDF["stopsequence"] == CurrentStop) & (RealTimeDF["a2eventtype"] == 0) & (RealTimeDF["platenumb"] == carID)
                    DepartResult = RealTimeDF[DepartQuery]["gpstime"].tolist()
                    DepartResult = [str(time) for time in DepartResult]

                    if not isValid(PreviousStopDepartTime, ArrivalResult, 1): ArrivalResult = []
                    if not isValid(PreviousStopDepartTime, DepartResult, 1): DepartResult = []

                    if len(ArrivalResult) != 0  and len(DepartResult) != 0:
                        ArrivalTime = ArrivalResult[FindClosestTime(PreviousStopDepartTime, ArrivalResult)]
                        DepartTime = DepartResult[FindClosestTime(PreviousStopDepartTime, DepartResult)]
                        CurrentStopStayTime = CalDiffTime(ArrivalTime, DepartTime)
                    else:
                        CurrentStopStayTime = 0

                """
                看當前站有沒有離站時間
                1. 如果沒有則代表沒有辦法跟其他站點計算GroundTruth(進站-離站)，直接跳過迴圈
                """
                DepartQuery = (RealTimeDF["routeid"] == ROUTEID) & (RealTimeDF["direction"] == DIRECTION) & (RealTimeDF["stopsequence"] == CurrentStop) & (RealTimeDF["a2eventtype"] == 0) & (RealTimeDF["platenumb"] == carID)
                DepartResult = RealTimeDF[DepartQuery]["gpstime"].tolist()
                DepartResult = [str(time) for time in DepartResult]

                if not isValid(PreviousStopDepartTime, DepartResult, 1): DepartResult = []

                if len(DepartResult) != 0:
                    CurrentStopDepartTime = DepartResult[FindClosestTime(PreviousStopDepartTime, DepartResult)]
                    PreviousStopDepartTime = CurrentStopDepartTime
                else:
                    continue
                
                """
                接下來要用一個迴圈去取得後面站點的進站時間(Ground Truth)、歷史行駛時間、API時間、歷史等待時間
                在第一個站點是不需要等待時間的計算的
                """
                ResultWS = wb[f"第{CurrentStop}站"]
                """
                維護兩個陣列，用來儲存累積的當前等待時間以及歷史等待時間
                """
                CurrentStayTimes, StastisticStayTimes = [], []

                for AfterStop in range(CurrentStop+1, len(RouteTable)+1):
                    Ratio = GetRatio(ROUTEID, DIRECTION, CurrentStop, carID, CurrentStopDepartTime, RealTimeDF, SechudleIndex, HistoryDriveTimeTable)
                    CurrentStayTime = sum(CurrentStayTimes)
                    HistoryStayTime = sum(StastisticStayTimes)
                    
                    HistoryDriveTime = GetStasticDriveTime(HistoryDriveTimeTable, CurrentStop, AfterStop, SechudleIndex)
                    GroundTruth = GetGroundTruth(ROUTEID, DIRECTION, AfterStop, carID, CurrentStopDepartTime, RealTimeDF)
                    KmeansCenter1, KmeansCenter2 = GetKmeans(KmeansGroup1, KmeansGroup2, CurrentStop, AfterStop, SechudleIndex)
                    NewDriveTime = sum(PredictedStatus(HistoryDriveTime, KmeansCenter1, KmeansCenter2, BusStatus, ROUTEID, model, device, DefaultValue, K))

                    """
                    計算完成上述資訊之後，需要將當前預測站的歷史等待時間以及真實等待時間加入到兩個串列之中進行累積
                    """
                    StastisticStayTimes.append(GetHistoryStayTime(ROUTEID, DAY, DIRECTION, AfterStop, SechudleIndex))
                    CurrentStayTimes.append(CurrentStopStayTime)
                    """
                    判斷groundTruth是否為離群直(1.5IQR)，如果是的話就跳過本次迴圈
                    """
                    UpperBound, LowerBound = GetIQR(ROUTEID, DAY, DIRECTION, CurrentStop, AfterStop, SechudleIndex)
                    if GroundTruth > UpperBound or GroundTruth < LowerBound:
                        continue

                    """
                    這裡要根據對應的站點，將資料寫入到對應的Excel活頁中 
                    """
                    if AfterStop - CurrentStop <= 2:
                        ResultWS.append([DATE, f"{SechudleIndex+1}-{AfterStop-CurrentStop}", sum(HistoryDriveTime), Ratio, HistoryStayTime, CurrentStayTime, sum(KmeansCenter1), sum(KmeansCenter2), NewDriveTime, GroundTruth])
                    else:
                        ResultWS.append([DATE, f"{SechudleIndex+1}-other", sum(HistoryDriveTime), Ratio, HistoryStayTime, CurrentStayTime, sum(KmeansCenter1), sum(KmeansCenter2), NewDriveTime, GroundTruth])
                
                """
                判斷經過的站點是屬於哪一群的
                """
                KmeansCenter1, KmeansCenter2 = GetKmeans(KmeansGroup1, KmeansGroup2, CurrentStop, CurrentStop+1, SechudleIndex)
                GroundTruth = GetGroundTruth(ROUTEID, DIRECTION, CurrentStop+1, carID, CurrentStopDepartTime, RealTimeDF)
                if abs(KmeansCenter1[0]-KmeansCenter2[0]) < K:
                    BusStatus.append(3)
                else:
                    if abs(KmeansCenter1[0]-GroundTruth) <= abs(KmeansCenter2[0]-GroundTruth):
                        BusStatus.append(1)
                    else:
                        BusStatus.append(2)

            print(f"日期:{DATE}, 第{SechudleIndex+1}班次整理完成") 
            wb.save(f"training_dataset/{ROUTEID}/{DIRECTION}/dataset_{DAY}.xlsx")

        print(f"{DATE} 儲存完成")